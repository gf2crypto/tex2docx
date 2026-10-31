"""Build input.docx: a chapter as a co-author would type it in Word.

Heading 1 + text with bold/italic, a bulleted and a numbered list, a 2x2 table,
an inline OMML formula and a display one (XML from fixtures/omml), a picture.
Run: python make.py  (python-docx; the OMML XML is pasted as raw elements).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.shared import Cm
from lxml import etree

HERE = Path(__file__).parent
OMML = HERE.parents[1] / "fixtures" / "omml"


def omml(name: str):
    return etree.fromstring((OMML / f"{name}.xml").read_bytes())


def main() -> None:
    doc = Document()
    doc.add_heading("Глава соавтора", level=1)
    p = doc.add_paragraph("Обычный текст с ")
    p.add_run("полужирным").bold = True
    p.add_run(" и ")
    p.add_run("курсивом").italic = True
    p.add_run(". Формула в строке: ")
    p._p.append(omml("inline-x2"))
    p.add_run(" — и дальше текст.")
    doc.add_paragraph("Первый пункт", style="List Bullet")
    doc.add_paragraph("Второй пункт", style="List Bullet")
    doc.add_paragraph("Шаг один", style="List Number")
    doc.add_paragraph("Шаг два", style="List Number")
    eq = doc.add_paragraph()
    eq.alignment = 1  # centered display formula typed in Word
    eq._p.append(omml("display-sum"))
    t = doc.add_table(rows=2, cols=2)
    t.style = "Table Grid"
    t.cell(0, 0).text, t.cell(0, 1).text = "a", "b"
    t.cell(1, 0).text, t.cell(1, 1).text = "1", "2"
    doc.add_picture(str(HERE.parents[1] / "fixtures" / "pic.png"), width=Cm(3))
    doc.add_paragraph("Подпись к рисунку").style = "Caption"
    doc.save(str(HERE / "input.docx"))


if __name__ == "__main__":
    main()
