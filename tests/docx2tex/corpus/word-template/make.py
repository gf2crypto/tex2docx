"""Build input.docx: a chapter typed in Word with the tex2docx template (reference.docx)
as the document template, but without bookmarks and without a companion file, as a
co-author would do it by hand.

Theorem / Proof / Equation / Algorithm (tabs for the nesting, no "end" lines) /
Bibliography paragraphs with "[1]" citations in the text; a heading typed with its number.
Run: python make.py  (python-docx; the OMML XML is pasted as raw elements).
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from lxml import etree

HERE = Path(__file__).parent
OMML = HERE.parents[1] / "fixtures" / "omml"
REFERENCE = HERE.parents[3] / "src" / "tex2docx" / "data" / "reference.docx"


def omml(name: str):
    return etree.fromstring((OMML / f"{name}.xml").read_bytes())


def style(doc, name: str):
    """python-docx does not find pandoc's 'heading 1' by the name 'Heading 1'."""
    for st in doc.styles:
        if st.name.lower() == name.lower() or st.style_id == name.replace(" ", ""):
            return st
    raise KeyError(name)


def main() -> None:
    doc = Document(str(REFERENCE))
    body = doc.element.body
    for el in list(body):  # the template's sample paragraphs
        if el.tag.endswith("}p") or el.tag.endswith("}tbl"):
            body.remove(el)
    doc.add_paragraph("2 Глава из Word", style=style(doc, "Heading 1"))
    p = doc.add_paragraph(style=style(doc, "Theorem"))
    p.add_run("Теорема 2.1").bold = True
    p.add_run(". ")
    p.add_run("Для любого ").italic = True
    p._p.append(omml("inline-x2"))
    p.add_run(" утверждение верно.").italic = True
    eq = doc.add_paragraph(style=style(doc, "Equation"))
    eq.add_run("\t")
    eq._p.append(omml("display-sum"))
    eq.add_run("\t(2.1)")
    p = doc.add_paragraph(style=style(doc, "Theorem"))  # the theorem goes on after its formula
    p.add_run("где ").italic = True
    p._p.append(omml("inline-x2"))
    p.add_run(" — квадрат.").italic = True
    p = doc.add_paragraph(style=style(doc, "Proof"))
    p.add_run("Доказательство.").italic = True
    p.add_run(" Следует из (2.1) и [1].\t◻")
    p = doc.add_paragraph(style=style(doc, "Algorithm Caption"))
    p.add_run("Алгоритм 1").bold = True
    p.add_run(". Поиск")
    lines = [("1:", 1, [("for", True), (" всех ", False), ("do", True)]),
             ("2:", 2, [("if", True), (" найдено ", False), ("then", True)]),
             ("3:", 3, [("return", True), (" истина", False)]),
             ("4:", 1, [("return", True), (" ложь", False)])]
    for num, level, parts in lines:
        p = doc.add_paragraph(style=style(doc, "Algorithm"))
        p.add_run(num + "\t" * level)
        for text, bold in parts:
            p.add_run(text).bold = bold
    doc.add_paragraph("Таблица 1. Значения", style=style(doc, "Table Caption"))
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "a", "b"
    t.cell(1, 0).text, t.cell(1, 1).text = "1", "2"
    doc.add_paragraph("См. также [1, 2].", style=style(doc, "Body Text"))
    doc.add_paragraph("Список литературы", style=style(doc, "Heading 1"))
    doc.add_paragraph("1. Первый источник.", style=style(doc, "Bibliography"))
    doc.add_paragraph("2. Второй источник.", style=style(doc, "Bibliography"))
    doc.save(str(HERE / "input.docx"))


if __name__ == "__main__":
    main()
